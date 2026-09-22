from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = Path(
    "/Users/teddy.lin/github/junior-subjects/初中/成绩排名/物理/第一次考试/"
    "物理第一次考试错题分析报告.docx"
)
FONT = "PingFang SC"
BLACK = "000000"
NAVY = "1F4E79"
PALE_BLUE = "D9EAF7"
PALE_GRAY = "F3F6F8"
BORDER = "D9D9D9"
RED = "A61B1B"


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:{}".format(edge)
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_text(cell, text, size=9, bold=False, color=BLACK, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)
    set_cell_border(cell)


def add_heading(document, text, level=1):
    paragraph = document.add_paragraph()
    paragraph.style = document.styles[f"Heading {level}"]
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 7)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    set_run_font(run, size=14 if level == 1 else 11.5, bold=True, color=BLACK)
    return paragraph


def add_body(document, text, before=0, after=4):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.28
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5, color=BLACK)
    return paragraph


def add_label_body(document, label, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.28
    label_run = paragraph.add_run(label)
    set_run_font(label_run, size=10.5, bold=True, color=NAVY)
    body_run = paragraph.add_run(text)
    set_run_font(body_run, size=10.5, color=BLACK)
    return paragraph


def set_table_widths(table, widths_cm):
    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)


def build():
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.55)
    section.bottom_margin = Cm(1.55)
    section.left_margin = Cm(1.55)
    section.right_margin = Cm(1.55)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.7)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    for name in ("Heading 1", "Heading 2"):
        styles[name].font.name = FONT
        styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        styles[name].font.color.rgb = RGBColor(0, 0, 0)
    title_style_ppr = styles["Title"]._element.find(qn("w:pPr"))
    if title_style_ppr is not None:
        title_border = title_style_ppr.find(qn("w:pBdr"))
        if title_border is not None:
            title_style_ppr.remove(title_border)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(5)
    title_run = title.add_run("物理第一次考试错题分析报告")
    set_run_font(title_run, size=20, bold=True, color=BLACK)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(12)
    subtitle_run = subtitle.add_run("2026-2027 学年第一学期  九年级物理阶段检测卷（一）")
    set_run_font(subtitle_run, size=10.5, color=BLACK)

    overview = document.add_table(rows=1, cols=3)
    overview.alignment = WD_TABLE_ALIGNMENT.CENTER
    overview.autofit = False
    values = [
        ("本次得分", "65 / 70 分"),
        ("得分率", "92.9%"),
        ("失分题目", "第 1、11、18（3）、18（5）题"),
    ]
    for cell, (label, value) in zip(overview.rows[0].cells, values):
        set_cell_shading(cell, PALE_BLUE)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        run = p.add_run(label)
        set_run_font(run, size=8.5, bold=True, color=NAVY)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(0)
        value_run = p2.add_run(value)
        set_run_font(value_run, size=12, bold=True, color=BLACK)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell, top=100, start=80, bottom=100, end=80)
        set_cell_border(cell)
    set_table_widths(overview, [5.0, 4.1, 7.4])

    conclusion = document.add_paragraph()
    conclusion.paragraph_format.space_before = Pt(9)
    conclusion.paragraph_format.space_after = Pt(8)
    conclusion.paragraph_format.line_spacing = 1.28
    lead = conclusion.add_run("总体判断：")
    set_run_font(lead, size=10.5, bold=True, color=NAVY)
    text = conclusion.add_run(
        "基础概念、实验探究和热量效率计算掌握较好；失分主要来自把物理规律放进真实情境时的判断，"
        "以及综合题中图示选择与文字理由没有完整作答。"
    )
    set_run_font(text, size=10.5, color=BLACK)

    add_heading(document, "错题逐题订正", 1)
    table = document.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    headers = ["题号", "学生作答", "正确答案", "失分核心"]
    for cell, text in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, NAVY)
        add_text(cell, text, size=9, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)

    wrong_rows = [
        (
            "1",
            "选 B",
            "选 D",
            "把温度差归因于汽化，忽略比热容。",
        ),
        (
            "11",
            "近地面箭头方向相反",
            "大海 → 陆地",
            "未从陆地受热更快、空气上升推导气流方向。",
        ),
        (
            "18（3）",
            "选 A",
            "选 D",
            "未依据液膜表面张力和“面积尽量减小”判断绳形。",
        ),
        (
            "18（5）",
            "未作答",
            "材料 B 不被油润湿",
            "理由性表述缺失，造成会而失分。",
        ),
    ]
    for index, row_values in enumerate(wrong_rows):
        row = table.add_row()
        for cell, text in zip(row.cells, row_values):
            if index % 2 == 1:
                set_cell_shading(cell, PALE_GRAY)
            add_text(
                cell,
                text,
                size=8.8,
                bold=(cell == row.cells[0]),
                color=BLACK,
                align=WD_ALIGN_PARAGRAPH.CENTER if cell in row.cells[:3] else WD_ALIGN_PARAGRAPH.LEFT,
            )
    set_table_widths(table, [1.2, 3.4, 3.1, 8.8])

    add_heading(document, "第 1 题  比热容的情境判断", 2)
    add_label_body(
        document,
        "正确答案：",
        "D，水的比热容比砂石的比热容大。",
    )
    add_label_body(
        document,
        "为什么：",
        "在相同日照条件下，水和砂石吸收相近的热量，水由于比热容较大，温度升高得较慢，"
        "因此水边或有水的地方温度相对较低。B 说“砂石不会汽化”本身也不成立，不能作为解释。",
    )
    add_label_body(
        document,
        "解题提示：",
        "看到“同样受热后谁升温更少”，优先联想 Q = cmΔt；在质量和吸热量相近时，c 越大，Δt 越小。",
    )

    add_heading(document, "第 11 题  海陆风作图", 2)
    add_label_body(document, "正确答案：", "炎热的白天，近地面气流方向应画为“大海 → 陆地”。")
    add_label_body(
        document,
        "为什么：",
        "白天陆地升温比海水快，陆地上方空气受热上升；海面附近较冷、较密的空气在近地面补充到陆地，"
        "形成海风。作图时先判断“哪里更热、哪里上升”，再画近地面补气方向。",
    )

    add_heading(document, "第 18（3）题  液膜与表面张力", 2)
    add_label_body(document, "正确答案：", "选 D。")
    add_label_body(
        document,
        "为什么：",
        "刺破 a 部分后，只剩 b 部分液膜。液膜受表面张力作用会尽量减小面积，"
        "因此棉线会向 b 区液膜一侧弯曲，呈现 D 图所示的形状。",
    )

    add_heading(document, "第 18（5）题  浸润与不浸润", 2)
    add_label_body(
        document,
        "规范答案：",
        "材料 B 不被油润湿（油与材料 B 间的相互作用力小于油分子之间的相互作用力）。",
    )
    add_label_body(
        document,
        "为什么：",
        "瓶口若用 B 材料，油不易在其表面铺展和附着，能减少沿瓶口流到瓶外的情况。"
        "这类题的关键不是只写“液滴更圆”，而是写出“不浸润”和分子间作用力大小关系。",
    )

    add_heading(document, "薄弱点诊断", 1)

    diagnosis = [
        (
            "一  从生活现象回到物理量",
            "第 1 题说明，对“水边更凉”这类生活情境，容易先抓住表面现象而没有回到比热容。"
            "要把条件翻译成公式语言：比较的是升温多少，而不是只判断是否发生汽化。",
        ),
        (
            "二  先讲机制，再画方向或选图",
            "第 11 题和第 18（3）题都属于图示推理。建议固定使用“条件 → 物理机制 → 结果”三步："
            "先写出受热上升或表面张力，再确定气流方向或液膜边界的移动方向。",
        ),
        (
            "三  分子间作用力的应用表达",
            "浸润、不浸润、表面张力的基础知识并非完全不会，但第 18（5）题没有落笔，"
            "说明需要练习把“看到的现象”转换为“相互作用力大小关系”的完整句子。",
        ),
        (
            "四  避免空题",
            "本次第 18（5）题属于可由材料直接推出的填空。考场最后应留 2 分钟专门检查空格，"
            "哪怕答案还不够完整，也先写出核心词，避免零分。",
        ),
    ]
    for heading, content in diagnosis:
        add_heading(document, heading, 2)
        add_body(document, content, after=4)

    add_heading(document, "已经掌握得较好的部分", 1)
    strength_table = document.add_table(rows=1, cols=2)
    strength_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    strength_table.autofit = False
    for cell, text in zip(strength_table.rows[0].cells, ["知识或能力", "本次表现"]):
        set_cell_shading(cell, NAVY)
        add_text(cell, text, size=9, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
    strengths = [
        ("热量、热值、效率计算", "第 10、16、17、19 题计算过程和结果正确，公式使用较稳。"),
        ("实验探究与控制变量", "第 13-15 题对扩散、内能变化、控制变量和转换法的回答较完整。"),
        ("能量转化与发动机基础", "第 5-10、12 题的能量转化、热机工作和效率判断总体正确。"),
        ("分子运动基础", "扩散、分子永不停息做无规则运动、分子间引力等基础结论掌握到位。"),
    ]
    for index, (topic, performance) in enumerate(strengths):
        row = strength_table.add_row()
        for cell, text in zip(row.cells, (topic, performance)):
            if index % 2 == 1:
                set_cell_shading(cell, PALE_GRAY)
            add_text(cell, text, size=9.2, bold=(cell == row.cells[0]), color=BLACK)
    set_table_widths(strength_table, [4.4, 12.1])

    add_heading(document, "针对性巩固安排", 1)
    plan = [
        (
            "第一轮 15 分钟",
            "复习比热容：做 3 道“相同质量、相同吸热或相同加热时间”的比较题。每题都写出谁的 Δt 更大及理由。",
        ),
        (
            "第二轮 20 分钟",
            "整理“表面张力、浸润、不浸润”三张小卡：现象、受力或作用力关系、典型例子各写一条。",
        ),
        (
            "第三轮 15 分钟",
            "专练图示题：海陆风、热传递方向、液膜绳形各画一遍，并在图旁写出依据。",
        ),
        (
            "考前 3 分钟",
            "按“选择题涂卡、填空是否有空、计算单位和百分号、作图箭头方向”四项检查答题卡。",
        ),
    ]
    for title, content in plan:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(5)
        paragraph.paragraph_format.line_spacing = 1.28
        prefix = paragraph.add_run(f"{title}：")
        set_run_font(prefix, size=10.5, bold=True, color=NAVY)
        detail = paragraph.add_run(content)
        set_run_font(detail, size=10.5, color=BLACK)

    add_heading(document, "复盘后自测", 1)
    add_body(
        document,
        "不看上面的订正答案，能把下列四项说清楚，说明这次失分点已经补上。",
        after=5,
    )
    check_table = document.add_table(rows=1, cols=3)
    check_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    check_table.autofit = False
    for cell, text in zip(check_table.rows[0].cells, ["检查点", "自测问题", "答案要点"]):
        set_cell_shading(cell, NAVY)
        add_text(cell, text, size=8.8, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
    checks = [
        (
            "比热容",
            "同质量的水和砂石吸收相近热量，谁升温更少？",
            "水。水的比热容较大，温度升高较少。",
        ),
        (
            "海陆风",
            "炎热白天，近地面空气从哪里流向哪里？",
            "大海流向陆地；陆地受热快，空气上升。",
        ),
        (
            "液膜",
            "a 区液膜刺破后，棉线形状选哪一图？",
            "选 D；剩余液膜在表面张力作用下尽量减小面积。",
        ),
        (
            "浸润",
            "为什么用材料 B 做油瓶瓶口更不易流油？",
            "油不润湿 B；油与 B 间的作用力较小。",
        ),
    ]
    for index, values in enumerate(checks):
        row = check_table.add_row()
        for cell, text in zip(row.cells, values):
            if index % 2 == 1:
                set_cell_shading(cell, PALE_GRAY)
            add_text(cell, text, size=8.7, bold=(cell == row.cells[0]), color=BLACK)
    set_table_widths(check_table, [3.0, 6.5, 7.0])

    document.core_properties.title = "物理第一次考试错题分析报告"
    document.core_properties.subject = "九年级物理错题订正与薄弱点分析"
    document.core_properties.author = "OpenAI Codex"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
